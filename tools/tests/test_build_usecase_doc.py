from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

import yaml
from docx import Document

from tools.build_usecase_doc import build_document


class BuildUsecaseDocTest(unittest.TestCase):
    @staticmethod
    def _create_template(path: Path) -> None:
        document = Document()
        document.add_heading("1. Danh sách 31 Use case", level=1)
        for code in "ABCDEFG":
            document.add_heading(f"{code}. Nhóm", level=2)
            rows = 2 if code == "A" else 1
            table = document.add_table(rows=rows, cols=3)
            for cell, value in zip(table.rows[0].cells, ("ID", "Tên Use Case", "Phạm vi")):
                cell.text = value
        document.add_heading("Quy ước thuật ngữ", level=1)
        document.add_paragraph(
            "Tên tiếng Việt được dùng trước; thuật ngữ tiếng Anh trong ngoặc giúp đối chiếu."
        )
        glossary = document.add_table(rows=2, cols=2)
        glossary.rows[0].cells[0].text = "Thuật ngữ trong tài liệu"
        glossary.rows[0].cells[1].text = "Ý nghĩa"
        document.add_heading("2. Use case chi tiết", level=1)
        document.add_heading("UC01: Ví dụ", level=2)
        document.add_table(rows=12, cols=4)
        path.parent.mkdir(parents=True)
        document.save(path)

    def test_build_is_deterministic_and_contains_sources(self) -> None:
        temp_root = Path.cwd() / ".tmp"
        temp_root.mkdir(exist_ok=True)
        with tempfile.TemporaryDirectory(dir=temp_root) as directory:
            root = Path(directory)
            sources = root / "plato_docs/usecases"
            sources.mkdir(parents=True)
            sections = (
                "Mục tiêu",
                "Tác nhân",
                "Tiền điều kiện",
                "Kích hoạt",
                "Luồng chính",
                "Luồng thay thế và ngoại lệ",
                "Hậu điều kiện",
                "Quy tắc nghiệp vụ",
                "Use Case liên quan",
                "Thuật ngữ",
            )
            body = "\n\n".join(f"## {section}\nNội dung {section}" for section in sections)
            (sources / "UC01.md").write_text(
                f"# UC01 — Ví dụ\n\n{body}\n", encoding="utf-8"
            )
            index = root / "plato_docs/usecase-index.yaml"
            index.write_text(
                yaml.safe_dump(
                    {
                        "usecases": {
                            "UC01": {
                                "title": "Ví dụ",
                                "source": "plato_docs/usecases/UC01.md",
                            }
                        }
                    },
                    allow_unicode=True,
                    sort_keys=False,
                ),
                encoding="utf-8",
            )
            template = root / "plato_docs/templates/usecase-template.docx"
            self._create_template(template)
            document_spec = {
                "version": 1,
                "template": "plato_docs/templates/usecase-template.docx",
                "glossary_note": "Tên tiếng Việt được dùng trước.",
                "groups": [
                    {
                        "title": f"{code}. Nhóm",
                        "usecases": ["UC01"] if code == "A" else [],
                    }
                    for code in "ABCDEFG"
                ],
                "glossary": [{"term": "Ví dụ", "definition": "Mô tả"}],
                "usecases": {
                    "UC01": {
                        "summary": "Tóm tắt",
                        "importance": "Cao",
                        "type": "essential",
                        "supporting_actor": "Không có",
                        "offstage_actor": "Không có",
                        "relationships": {
                            "entry_point": "Màn hình ví dụ",
                            "association": "Người dùng",
                            "include": "Không có",
                            "extend": "Không có",
                        },
                    }
                },
            }
            (root / "plato_docs/usecase-document.yaml").write_text(
                yaml.safe_dump(document_spec, allow_unicode=True, sort_keys=False),
                encoding="utf-8",
            )
            first = root / "first.docx"
            second = root / "second.docx"
            build_document(root, "plato_docs/usecase-index.yaml", first)
            build_document(root, "plato_docs/usecase-index.yaml", second)

            self.assertEqual(first.read_bytes(), second.read_bytes())
            generated = Document(first)
            text = "\n".join(p.text for p in generated.paragraphs)
            self.assertIn("UC01: Ví dụ", text)
            self.assertEqual(len(generated.tables), 9)
            self.assertEqual(len(generated.tables[8].rows), 14)
            self.assertIn("Nội dung Mục tiêu", generated.tables[8].rows[4].cells[1].text)


if __name__ == "__main__":
    unittest.main()
