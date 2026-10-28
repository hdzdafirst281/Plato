from __future__ import annotations

import argparse
import re
import tempfile
import zipfile
from copy import deepcopy
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

try:
    from tools.usecase_common import (
        DEFAULT_DOCUMENT_SPEC,
        DEFAULT_INDEX,
        load_document_spec,
        load_usecase_index,
        repository_path,
        repository_root,
    )
    from tools.validate_usecases import validate_repository
except ModuleNotFoundError:
    from usecase_common import (  # type: ignore[no-redef]
        DEFAULT_DOCUMENT_SPEC,
        DEFAULT_INDEX,
        load_document_spec,
        load_usecase_index,
        repository_path,
        repository_root,
    )
    from validate_usecases import validate_repository  # type: ignore[no-redef]


DEFAULT_OUTPUT = Path("plato_docs/generated/plato_usecases.docx")
FIXED_TIMESTAMP = datetime(2026, 10, 27, tzinfo=timezone.utc)
ZIP_TIMESTAMP = (2026, 10, 27, 0, 0, 0)
HEADING_PATTERN = re.compile(r"^#\s+(UC\d{2})\s+[—–-]\s+(.+?)\s*$", re.MULTILINE)
SECTION_PATTERN = re.compile(r"^##\s+(.+?)\s*$", re.MULTILINE)
INLINE_PATTERN = re.compile(r"(\*\*.+?\*\*|`.+?`)")
FLOW_HEADING_PATTERN = re.compile(r"^\*\*([SAE]\d+)\s*[—–-].+\*\*$")
DETAIL_LABELS = (
    "Use Case Name",
    "Use Case Type",
    "Primary Actor",
    "Supporting Actor and Offstage Actor",
    "Brief Description",
    "Trigger",
    "Pre-conditions",
    "Post-conditions",
    "Relationships",
    "Normal Flow of Events",
    "Subflows",
    "Alternate/Exceptional Flows",
    "Business Rules",
    "Terminology",
)


def _font(run, name: str = "Times New Roman", size: float = 12) -> None:
    run.font.name = name
    fonts = run._element.get_or_add_rPr().rFonts
    fonts.set(qn("w:ascii"), name)
    fonts.set(qn("w:hAnsi"), name)
    fonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)


def _clear_paragraph(paragraph) -> None:
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def _prepare_cell(cell):
    paragraph = cell.paragraphs[0]
    _clear_paragraph(paragraph)
    for extra in list(cell.paragraphs[1:]):
        cell._tc.remove(extra._p)
    return paragraph


def _add_inline(paragraph, text: str, *, default_bold: bool = False) -> None:
    for part in INLINE_PATTERN.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
            _font(run)
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            _font(run, "Courier New")
        else:
            run = paragraph.add_run(part)
            run.bold = default_bold
            _font(run)


def _write_text(cell, text: str, *, bold: bool = False) -> None:
    paragraph = _prepare_cell(cell)
    lines = text.strip().splitlines() if text.strip() else [""]
    for position, line in enumerate(lines):
        if position:
            paragraph.add_run().add_break()
        _add_inline(paragraph, line, default_bold=bold)


def _write_labeled_value(cell, label: str, value: str) -> None:
    paragraph = _prepare_cell(cell)
    run = paragraph.add_run(f"{label}: ")
    run.bold = True
    _font(run)
    _add_inline(paragraph, value)


def _write_labeled_blocks(cell, blocks: list[tuple[str, str]]) -> None:
    paragraph = _prepare_cell(cell)
    for position, (label, value) in enumerate(blocks):
        if position:
            paragraph.add_run().add_break()
            paragraph.add_run().add_break()
        run = paragraph.add_run(f"{label}:")
        run.bold = True
        _font(run)
        if value:
            paragraph.add_run().add_break()
            for line_position, line in enumerate(value.strip().splitlines()):
                if line_position:
                    paragraph.add_run().add_break()
                _add_inline(paragraph, line)


def _write_usecase_type(cell, value: str) -> None:
    paragraph = _prepare_cell(cell)
    selections = (
        ("☑" if value == "essential" else "☐", " Essential "),
        ("☑" if value == "temporal" else "☐", " Temporal"),
    )
    for symbol, label in selections:
        run = paragraph.add_run(symbol)
        _font(run, "Segoe UI Symbol")
        run = paragraph.add_run(label)
        _font(run)


def _write_heading(paragraph, text: str) -> None:
    _clear_paragraph(paragraph)
    run = paragraph.add_run(text)
    run.bold = True
    _font(run, size=18)


def _parse_markdown(markdown: str, source: Path) -> tuple[str, str, dict[str, str]]:
    heading = HEADING_PATTERN.search(markdown)
    if not heading:
        raise ValueError(f"Invalid source heading: {source}")
    sections: dict[str, str] = {}
    matches = list(SECTION_PATTERN.finditer(markdown))
    for position, match in enumerate(matches):
        start = match.end()
        end = matches[position + 1].start() if position + 1 < len(matches) else len(markdown)
        sections[match.group(1)] = markdown[start:end].strip()
    return heading.group(1), heading.group(2), sections


def _split_variant_flows(value: str) -> tuple[str, str]:
    blocks = [block.strip() for block in re.split(r"\n\s*\n", value) if block.strip()]
    subflows: list[str] = []
    alternatives: list[str] = []
    current = alternatives
    for block in blocks:
        first_line = block.splitlines()[0]
        match = FLOW_HEADING_PATTERN.match(first_line)
        if match:
            current = subflows if match.group(1).startswith("S") else alternatives
        current.append(block)
    return "\n\n".join(subflows) or "Không có", "\n\n".join(alternatives) or "Không có"


def _append_detail_row(table) -> None:
    table._tbl.append(deepcopy(table.rows[-1]._tr))


def _ensure_template_shape(document: Document, spec: dict[str, Any]) -> None:
    groups = spec["groups"]
    glossary = spec["glossary"]
    metadata = spec["usecases"]
    if len(document.tables) != 8 + len(metadata):
        raise ValueError(
            f"Template must contain {8 + len(metadata)} tables; found {len(document.tables)}"
        )
    for position, group in enumerate(groups):
        expected_rows = len(group["usecases"]) + 1
        if len(document.tables[position].rows) != expected_rows:
            raise ValueError(
                f"Template group table {position + 1} must contain {expected_rows} rows"
            )
    if len(document.tables[7].rows) != len(glossary) + 1:
        raise ValueError("Template glossary row count differs from document spec")
    for position, table in enumerate(document.tables[8:], start=1):
        if len(table.rows) not in {12, 14} or len(table.columns) != 4:
            raise ValueError(f"Template detail table {position} has an unexpected shape")


def _populate_overview(
    document: Document,
    spec: dict[str, Any],
    usecases: dict[str, dict[str, Any]],
) -> None:
    group_headings = [
        paragraph
        for paragraph in document.paragraphs
        if paragraph.style.name == "Heading 2" and re.match(r"^[A-G]\.\s", paragraph.text)
    ]
    if len(group_headings) != len(spec["groups"]):
        raise ValueError("Template group heading count differs from document spec")
    for table, heading, group in zip(
        document.tables[:7], group_headings, spec["groups"], strict=True
    ):
        _write_heading(heading, group["title"])
        for row, uc_id in zip(table.rows[1:], group["usecases"], strict=True):
            _write_text(row.cells[0], uc_id, bold=True)
            _write_text(row.cells[1], usecases[uc_id]["title"])
            _write_text(row.cells[2], spec["usecases"][uc_id]["summary"])

    note = next(
        (
            paragraph
            for paragraph in document.paragraphs
            if paragraph.text.startswith("Tên tiếng Việt được dùng trước")
        ),
        None,
    )
    if note is None:
        raise ValueError("Template glossary note was not found")
    _clear_paragraph(note)
    _add_inline(note, spec["glossary_note"])

    glossary_table = document.tables[7]
    for row, item in zip(glossary_table.rows[1:], spec["glossary"], strict=True):
        _write_text(row.cells[0], item["term"])
        _write_text(row.cells[1], item["definition"])


def _populate_detail_table(
    table,
    uc_id: str,
    title: str,
    sections: dict[str, str],
    metadata: dict[str, Any],
) -> None:
    while len(table.rows) < len(DETAIL_LABELS):
        _append_detail_row(table)
    for row, label in zip(table.rows, DETAIL_LABELS, strict=True):
        _write_text(row.cells[0], label, bold=True)

    _write_labeled_value(table.rows[0].cells[0], "Use Case Name", title)
    _write_labeled_value(table.rows[0].cells[2], "ID", uc_id)
    _write_labeled_value(table.rows[0].cells[3], "Importance Level", metadata["importance"])
    _write_usecase_type(table.rows[1].cells[1], metadata["type"])
    primary_actor = next(
        line.strip() for line in sections["Tác nhân"].splitlines() if line.strip()
    )
    _write_text(table.rows[2].cells[1], primary_actor)
    _write_labeled_blocks(
        table.rows[3].cells[1],
        [
            ("Supporting Actor", metadata["supporting_actor"]),
            ("Offstage Actor", metadata["offstage_actor"]),
        ],
    )
    _write_text(table.rows[4].cells[1], sections["Mục tiêu"])
    _write_text(table.rows[5].cells[1], sections["Kích hoạt"])
    _write_text(table.rows[6].cells[1], sections["Tiền điều kiện"])
    _write_text(table.rows[7].cells[1], sections["Hậu điều kiện"])
    relationships = metadata["relationships"]
    _write_labeled_blocks(
        table.rows[8].cells[1],
        [
            ("Entry Point", relationships["entry_point"]),
            ("Association", relationships["association"]),
            ("Include", relationships["include"]),
            ("Extend", relationships["extend"]),
            ("Related Use Cases", sections["Use Case liên quan"]),
        ],
    )
    _write_text(table.rows[9].cells[1], sections["Luồng chính"])
    subflows, alternatives = _split_variant_flows(sections["Luồng thay thế và ngoại lệ"])
    _write_text(table.rows[10].cells[1], subflows)
    _write_text(table.rows[11].cells[1], alternatives)
    _write_text(table.rows[12].cells[1], sections["Quy tắc nghiệp vụ"])
    _write_text(table.rows[13].cells[1], sections["Thuật ngữ"])


def _normalize_docx(path: Path) -> None:
    with zipfile.ZipFile(path, "r") as source:
        entries = [(item.filename, source.read(item.filename)) for item in source.infolist()]
    with tempfile.NamedTemporaryFile(dir=path.parent, suffix=".docx", delete=False) as stream:
        normalized = Path(stream.name)
    try:
        with zipfile.ZipFile(normalized, "w", compression=zipfile.ZIP_DEFLATED) as target:
            for name, data in sorted(entries):
                info = zipfile.ZipInfo(name, ZIP_TIMESTAMP)
                info.compress_type = zipfile.ZIP_DEFLATED
                info.external_attr = 0o600 << 16
                target.writestr(info, data)
        normalized.replace(path)
    finally:
        normalized.unlink(missing_ok=True)


def build_document(
    root: Path,
    index_path: str | Path,
    output: Path,
    document_spec_path: str | Path = DEFAULT_DOCUMENT_SPEC,
    template_path: str | Path | None = None,
) -> None:
    _, usecases = load_usecase_index(root, index_path)
    _, spec = load_document_spec(root, document_spec_path)
    template = repository_path(root, template_path or spec["template"])
    document = Document(template)
    _ensure_template_shape(document, spec)
    _populate_overview(document, spec, usecases)

    uc_headings = {
        match.group(1): paragraph
        for paragraph in document.paragraphs
        if (match := re.match(r"^(UC\d{2}):", paragraph.text))
    }
    ordered_ids = [uc_id for group in spec["groups"] for uc_id in group["usecases"]]
    if set(uc_headings) != set(ordered_ids):
        raise ValueError("Template UC headings differ from document spec")
    for position, uc_id in enumerate(ordered_ids):
        entry = usecases[uc_id]
        source = repository_path(root, entry["source"])
        markdown = source.read_text(encoding="utf-8")
        heading_id, heading_title, sections = _parse_markdown(markdown, source)
        if heading_id != uc_id or heading_title != entry["title"]:
            raise ValueError(f"Source heading differs from index: {source}")
        _write_heading(uc_headings[uc_id], f"{uc_id}: {entry['title']}")
        _populate_detail_table(
            document.tables[8 + position],
            uc_id,
            entry["title"],
            sections,
            spec["usecases"][uc_id],
        )

    document.core_properties.title = "Đặc tả Use Case Plato"
    document.core_properties.subject = "Đặc tả 31 Use Case của ứng dụng Plato"
    document.core_properties.author = "Plato"
    document.core_properties.last_modified_by = "Plato Docs Generator"
    document.core_properties.created = FIXED_TIMESTAMP
    document.core_properties.modified = FIXED_TIMESTAMP
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    _normalize_docx(output)


def main() -> int:
    parser = argparse.ArgumentParser(description="Build the Plato Use Case DOCX")
    parser.add_argument("--root", help="Repository root")
    parser.add_argument("--index", default=str(DEFAULT_INDEX), help="Index path")
    parser.add_argument(
        "--document-spec", default=str(DEFAULT_DOCUMENT_SPEC), help="Document metadata path"
    )
    parser.add_argument("--template", help="Override the DOCX template path")
    parser.add_argument("--output", default=str(DEFAULT_OUTPUT), help="Output DOCX")
    parser.add_argument("--check", action="store_true", help="Fail when generated output is stale")
    args = parser.parse_args()

    root = repository_root(args.root)
    issues = validate_repository(root, args.index, document_spec_path=args.document_spec)
    if issues:
        for issue in issues:
            print(f"[{issue.level}] {issue.message}")
        return 1

    output = repository_path(root, args.output)
    if args.check:
        if not output.is_file():
            print(f"Generated DOCX is missing: {output.relative_to(root)}")
            return 1
        (root / ".tmp").mkdir(exist_ok=True)
        with tempfile.TemporaryDirectory(dir=root / ".tmp") as directory:
            candidate = Path(directory) / output.name
            build_document(
                root,
                args.index,
                candidate,
                args.document_spec,
                args.template,
            )
            if candidate.read_bytes() != output.read_bytes():
                print(f"Generated DOCX is stale: {output.relative_to(root)}")
                return 1
        print(f"Generated DOCX is current: {output.relative_to(root)}")
        return 0

    build_document(root, args.index, output, args.document_spec, args.template)
    print(f"Generated DOCX: {output.relative_to(root)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
